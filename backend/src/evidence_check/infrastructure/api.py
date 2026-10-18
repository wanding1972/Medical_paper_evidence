"""IO boundaries: immutable policy loading, local documents and safe rendering."""
from html import escape
from io import BytesIO
import json
import logging
from pathlib import Path
import re
from zipfile import BadZipFile, ZipFile
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import nsmap
from lxml.etree import XMLSyntaxError, XPathError, XPath
from evidence_check.contracts.errors import DocumentReadError, PolicyError
from evidence_check.contracts.models import AuditReport, FileFacts
from evidence_check.contracts.policy import Policy
from evidence_check.contracts.predicates import enforce

_LOG = logging.getLogger(__name__)


def _unique_object(pairs: list[tuple[str, object]]) -> dict:
    data = dict(pairs)
    if len(data) != len(pairs):
        raise ValueError("Duplicate object keys")
    return data


def load_policy(config_dir: Path) -> Policy:
    files = {"audit":"audit_rules.json", "limits":"limits.json", "parsing":"parsing.json",
             "retrieval":"retrieval.json", "acceptance":"acceptance.json"}
    try:
        data = {name:json.loads((config_dir / file).read_text(encoding="utf-8"), object_pairs_hook=_unique_object)
                for name, file in files.items()}
        policy = Policy.model_validate_json(json.dumps(data))
        XPath(policy.parsing.docx_unsupported_xpath, namespaces=nsmap)
        return policy
    except (OSError, UnicodeError, ValueError, re.error, XPathError) as error:
        _LOG.error("policy_invalid")
        raise PolicyError("policy_invalid", "Cannot load a complete, valid policy configuration.") from error


def _file_facts(path: Path, expanded_bytes: int = 0, unsupported_structure_count: int = 0) -> FileFacts:
    return FileFacts(suffix=path.suffix.lower(), size_bytes=path.stat().st_size,
                     expanded_bytes=expanded_bytes, unsupported_structure_count=unsupported_structure_count)


def _read_docx(path: Path, policy: Policy) -> str:
    with ZipFile(path) as archive:
        size = sum(info.file_size for info in archive.infolist())
        enforce(policy.limits.file_gates, _file_facts(path, size))
    document = Document(str(path))
    structures = document.element.xpath(policy.parsing.docx_unsupported_xpath)
    enforce(policy.limits.file_gates, _file_facts(path, size, len(structures)))
    return policy.parsing.paragraph_separator.join(paragraph.text for paragraph in document.paragraphs)


def read_document(path: Path, policy: Policy) -> str:
    try:
        facts = _file_facts(path)
        enforce(policy.limits.file_gates, facts)
        readers = {suffix:lambda: path.read_text(encoding="utf-8") for suffix in policy.limits.formats.markdown}
        readers.update({suffix:lambda: _read_docx(path, policy) for suffix in policy.limits.formats.docx})
        reader = readers.get(facts.suffix)
        if reader is None:
            raise DocumentReadError("format_adapter_missing", "No reader registered for this format.")
        text = reader()
        return text.removeprefix("\ufeff") if policy.parsing.strip_bom else text
    except (OSError, UnicodeError, BadZipFile, PackageNotFoundError, XMLSyntaxError, XPathError, KeyError, ValueError) as error:
        _LOG.error("document_read_failed")
        raise DocumentReadError("document_read_failed", "Cannot read the submitted document.") from error


def read_context(path: Path) -> dict:
    try:
        data = json.loads(path.read_text(encoding="utf-8"), object_pairs_hook=_unique_object)
        if not isinstance(data, dict) or "text" in data:
            raise ValueError("Context must be an object without text")
        return data
    except (OSError, UnicodeError, ValueError) as error:
        _LOG.error("context_read_failed")
        raise DocumentReadError("context_read_failed", "Cannot read a valid context object.") from error


def _safe(text: str) -> str:
    encoded = escape(text)
    return re.sub(r"([\\`*_{}\[\]()#+.!|>-])", r"\\\1", encoded).replace("\n", " ").replace("\r", " ")


def render_markdown(report: AuditReport) -> str:
    lines = ["# Evidence Check", "", _safe(report.audit_scope), "", _safe(report.disclaimer), "",
             f"Input SHA256: {report.input_sha256}", f"Policy: {_safe(report.policy_version)}", "",
             f"Policy SHA256: {report.policy_sha256}", "",
             f"Manual claims: {report.summary.manual_claim_count}",
             f"Findings: {report.summary.finding_count}", ""]
    for finding in report.findings:
        where = finding.location
        lines.extend([
            f"## {_safe(finding.rule_id)} / {_safe(finding.severity)} / {_safe(finding.review_status)}",
            f"Label: {_safe(finding.label)}",
            f"Paragraph {where.paragraph_index}, sentence {where.sentence_index}, chars [{where.char_start}, {where.char_end})",
            f"References: {', '.join(str(n) for n in finding.citation_ids)}",
            _safe(finding.claim_text), _safe(finding.message), _safe(finding.guidance), "",
        ])
    return "\n".join(lines)


def load_product_policy(config_dir: Path):
    from evidence_check.contracts.product_policy import ProductPolicy
    try:
        data=json.loads((config_dir/'product.json').read_text(encoding='utf-8'),object_pairs_hook=_unique_object)
        return ProductPolicy.model_validate_json(json.dumps(data))
    except (OSError,UnicodeError,ValueError,re.error) as error:
        _LOG.error('product_policy_invalid')
        raise PolicyError('product_policy_invalid','Cannot load a valid product policy.') from error


def read_uploaded_document(data: bytes,suffix: str,policy: Policy) -> str:
    """Parse uploaded bytes without leaving decrypted temporary files on disk."""
    try:
        facts=FileFacts(suffix=suffix,size_bytes=len(data),expanded_bytes=0,unsupported_structure_count=0)
        enforce(policy.limits.file_gates,facts)
        if suffix in policy.limits.formats.markdown:
            text=data.decode('utf-8')
        elif suffix in policy.limits.formats.docx:
            with ZipFile(BytesIO(data)) as archive:size=sum(i.file_size for i in archive.infolist())
            enforce(policy.limits.file_gates,facts.model_copy(update={'expanded_bytes':size}))
            doc=Document(BytesIO(data))
            count=len(doc.element.xpath(policy.parsing.docx_unsupported_xpath))
            enforce(policy.limits.file_gates,facts.model_copy(update={'expanded_bytes':size,'unsupported_structure_count':count}))
            text=policy.parsing.paragraph_separator.join(p.text for p in doc.paragraphs)
        else:raise DocumentReadError('format_adapter_missing','No reader registered for this format.')
        return text.removeprefix('\ufeff') if policy.parsing.strip_bom else text
    except (OSError,UnicodeError,BadZipFile,PackageNotFoundError,XMLSyntaxError,XPathError,KeyError,ValueError) as error:
        _LOG.error('document_read_failed')
        raise DocumentReadError('document_read_failed','Cannot read the submitted document.') from error
