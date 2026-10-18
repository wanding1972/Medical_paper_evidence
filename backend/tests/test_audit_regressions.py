"""Stage-6 regressions discovered by boundary and configuration review."""
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from docx import Document
from evidence_check.application.api import audit_document
from evidence_check.contracts.errors import InputError, PolicyError
from evidence_check.domains.document.api import parse_document
from evidence_check.infrastructure.api import load_policy, read_document
from support import BACKEND, document, policy


class AuditRegressions(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        self.config=self.root/'config';shutil.copytree(BACKEND/'config',self.config)

    def change(self,file,change):
        path=self.config/file;data=json.loads(path.read_text());change(data);path.write_text(json.dumps(data))

    def test_giant_integer_is_stable_input_error(self):
        with self.assertRaises(InputError):
            parse_document(document('Synthetic['+'1'*5000+'].'),policy())

    def test_docx_table_is_not_silently_omitted(self):
        path=self.root/'table.docx';word=Document();word.add_paragraph('Synthetic[1].')
        word.add_table(rows=1,cols=1).cell(0,0).text='Hidden claim[9].';word.save(path)
        with self.assertRaises(InputError):read_document(path,policy())

    def test_policy_type_mismatch_is_rejected(self):
        self.change('limits.json',lambda d:d['document_gates'][0]['when'][0].update(value=0))
        with self.assertRaises(PolicyError):load_policy(self.config)

    def test_missing_validation_field_is_rejected(self):
        self.change('limits.json',lambda d:d.update(claim_gates=d['claim_gates'][1:]))
        with self.assertRaises(PolicyError):load_policy(self.config)

    def test_parser_cannot_silently_drop_input(self):
        self.change('parsing.json',lambda d:d.update(sentence_pattern='Synthetic'))
        with self.assertRaises(InputError):
            parse_document(document('Synthetic[9].'),load_policy(self.config))

    def test_duplicate_config_keys_are_rejected(self):
        path=self.config/'audit_rules.json'
        path.write_text(path.read_text().replace('"version": "0.1"','"version": "wrong", "version": "0.1"'))
        with self.assertRaises(PolicyError):load_policy(self.config)

    def test_report_tracks_policy_content_not_just_version(self):
        first=audit_document(document(),policy())
        self.change('audit_rules.json',lambda d:d.update(audit_scope='Different scope'))
        second=audit_document(document(),load_policy(self.config))
        self.assertNotEqual(first.policy_sha256,second.policy_sha256)
        self.assertEqual(first.input_sha256,second.input_sha256)

    def test_pii_in_reference_title_is_blocked(self):
        from evidence_check.contracts.models import Reference
        with self.assertRaises(InputError):
            parse_document(document(references=(Reference(reference_id=1,title='Patient name: PRIVATE_NAME'),)),policy())
